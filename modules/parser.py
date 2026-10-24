import io

def extract_text(uploaded_file):
    """Extract text from an uploaded Streamlit file. Supports PDF, DOCX and TXT."""
    if uploaded_file is None:
        return ""
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if name.endswith(".txt"):
        return raw.decode("utf-8", errors="replace")
    if name.endswith(".pdf"):
        import fitz
        doc = fitz.open(stream=raw, filetype="pdf")
        return "\n".join(page.get_text("text") for page in doc)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(raw))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(chunks)
    raise ValueError("Unsupported file type. Use PDF, DOCX, or TXT.")
