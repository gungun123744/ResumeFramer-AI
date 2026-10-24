from io import BytesIO
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
import re

def make_docx(text):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    for line in text.splitlines():
        line = line.strip()
        if not line: continue
        if line.startswith("# "): doc.add_heading(line[2:], level=1)
        elif line.startswith("## "): doc.add_heading(line[3:], level=2)
        elif line.startswith("### "): doc.add_heading(line[4:], level=3)
        elif line.startswith(("- ", "* ")): doc.add_paragraph(line[2:], style="List Bullet")
        else: doc.add_paragraph(re.sub(r"\*\*(.*?)\*\*", r"\1", line))
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()

def make_report_pdf(title, text):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=.65*inch, leftMargin=.65*inch,
                            topMargin=.6*inch, bottomMargin=.6*inch)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for line in text.splitlines():
        line = line.strip()
        if not line:
            story.append(Spacer(1, 5)); continue
        if line.startswith("# "): story.append(Paragraph(line[2:], styles["Heading1"]))
        elif line.startswith("## "): story.append(Paragraph(line[3:], styles["Heading2"]))
        elif line.startswith(("- ", "* ")): story.append(Paragraph("• " + line[2:], styles["BodyText"]))
        else:
            clean = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            clean = clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(clean, styles["BodyText"]))
        story.append(Spacer(1, 4))
    doc.build(story)
    return buf.getvalue()
